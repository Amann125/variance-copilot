"""
Variance Copilot (reference solution)
-------------------------------------
Joins budget and actuals, flags material variances, drafts commentary
with an AI model (or a simple template if no API key is set), checks
every number in the commentary against the data, and writes an Excel report.

Run from the variance-copilot folder:
    python solution/variance_copilot.py
    python solution/variance_copilot.py --month 2026-04
"""
import argparse
import json
import os
import re
from pathlib import Path

import pandas as pd

# ---------- Settings (the Python version of blue input cells) ----------
PCT_THRESHOLD = 0.05      # flag if at least 5% off budget
DOLLAR_THRESHOLD = 5000   # ...and at least $5,000 off budget
MODEL = "claude-sonnet-5-5"

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "output"


def load_and_join():
    budget = pd.read_csv(DATA / "budget.csv")
    actuals = pd.read_csv(DATA / "actuals.csv")
    notes = pd.read_csv(DATA / "notes.csv")
    df = actuals.merge(budget, on=["month", "dept", "account"], suffixes=("_actual", "_budget"))
    df = df.merge(notes, on=["month", "dept", "account"], how="left")
    df["note"] = df["note"].fillna("")
    df["variance"] = df["amount_actual"] - df["amount_budget"]
    df["variance_pct"] = df["variance"] / df["amount_budget"]
    return df


def direction(row):
    if row["account"] == "Revenue":
        return "Favorable" if row["variance"] >= 0 else "Unfavorable"
    return "Unfavorable" if row["variance"] > 0 else "Favorable"


def flag(df):
    df = df.copy()
    df["direction"] = df.apply(direction, axis=1)
    df["material"] = (df["variance_pct"].abs() >= PCT_THRESHOLD) & (df["variance"].abs() >= DOLLAR_THRESHOLD)
    return df[df["material"]].sort_values("variance", key=abs, ascending=False)


def template_commentary(r):
    driver = r["note"] if r["note"] else "driver to confirm"
    return (f"{r['dept']} {r['account']} was ${abs(r['variance']):,.0f} ({abs(r['variance_pct']) * 100:.0f}%) "
            f"{r['direction'].lower()} to budget, driven by: {driver}.")


def ai_commentary(flags):
    """Returns a list of dicts: dept, account, commentary. Uses the API if a key is set."""
    if not os.environ.get("ANTHROPIC_API_KEY"):
        print("No ANTHROPIC_API_KEY found, so using template commentary. (That's fine. Add a key in session 6.)")
        return [{"dept": r["dept"], "account": r["account"], "commentary": template_commentary(r)} for _, r in flags.iterrows()]

    import anthropic
    client = anthropic.Anthropic()
    rows = []
    for _, r in flags.iterrows():
        rows.append({"dept": r["dept"], "account": r["account"], "budget": int(r["amount_budget"]),
                     "actual": int(r["amount_actual"]), "variance": int(r["variance"]),
                     "variance_pct": round(float(r["variance_pct"]) * 100), "direction": r["direction"], "note": r["note"]})
    prompt = f"""You are an FP&A analyst writing monthly variance commentary.
For each row, write ONE sentence: the variance in dollars and percent, favorable or unfavorable,
and the driver from the note. Use ONLY the numbers and notes given. Do not calculate new numbers.
If the note is blank, say "driver to confirm".
Return ONLY a JSON list of objects with keys "dept", "account", "commentary". No other text.

Rows:
{json.dumps(rows, indent=2)}"""
    msg = client.messages.create(model=MODEL, max_tokens=1500, messages=[{"role": "user", "content": prompt}])
    text = msg.content[0].text.strip().replace("```json", "").replace("```", "").strip()
    return json.loads(text)


def numbers_in(text):
    """Dollar amounts ($12,345) and percents (14%) that appear in the text."""
    found = re.findall(r"\$\s?([\d,]+(?:\.\d+)?)\s?(K|k|M)?|([\d.]+)\s?%", text)
    nums = []
    for dollars, unit, pct in found:
        if dollars:
            nums.append(float(dollars.replace(",", "")))
        elif pct:
            nums.append(float(pct))
    return nums


def verify(flags, results):
    allowed = set()
    for _, r in flags.iterrows():
        for v in (r["amount_budget"], r["amount_actual"], r["variance"]):
            allowed.add(round(abs(v)))
            allowed.add(round(abs(v) / 1000))
            allowed.add(round(abs(v) / 1000, 1))
        allowed.add(round(abs(r["variance_pct"]) * 100))
    for item in results:
        bad = [n for n in numbers_in(item["commentary"]) if round(n) not in allowed and round(n, 1) not in allowed]
        item["check"] = "OK" if not bad else f"Review: {bad}"
    return results


def write_report(df, flags, results, month):
    OUT.mkdir(exist_ok=True)
    summary = (df[df["account"] != "Revenue"].groupby("dept")[["amount_budget", "amount_actual", "variance"]]
               .sum().reset_index().rename(columns={"amount_budget": "budget", "amount_actual": "actual"}))
    out = flags.merge(pd.DataFrame(results), on=["dept", "account"], how="left")
    out = out[["month", "dept", "account", "amount_budget", "amount_actual", "variance", "variance_pct", "direction", "commentary", "check"]]
    path = OUT / f"variance_report_{month}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Expense summary", index=False)
        out.to_excel(xw, sheet_name="Flagged variances", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(name="Calibri", bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3A6E")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 16
        fv = xw.book["Flagged variances"]
        fv.column_dimensions["I"].width = 80
        for row in fv.iter_rows(min_row=2):
            for c in row[3:6]: c.number_format = "$#,##0;($#,##0)"
            row[6].number_format = "0.0%"
            if row[7].value == "Unfavorable":
                row[7].font = Font(name="Calibri", color="B42318", bold=True)
        for row in xw.book["Expense summary"].iter_rows(min_row=2):
            for c in row[1:4]: c.number_format = "$#,##0;($#,##0)"
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", default="2026-06")
    args = ap.parse_args()

    df = load_and_join()
    month_df = df[df["month"] == args.month]
    flags = flag(month_df)
    print(f"\n{len(flags)} material variance(s) in {args.month}:\n")
    print(flags[["dept", "account", "variance", "variance_pct", "direction"]].to_string(index=False))

    if flags.empty:
        print("\nNothing crossed the thresholds this month. Try another month with --month.")
        return
    results = verify(flags, ai_commentary(flags))
    print("\nCommentary:")
    for item in results:
        print(f" - [{item['check']}] {item['commentary']}")

    path = write_report(month_df, flags, results, args.month)
    print(f"\nReport saved: {path}")


if __name__ == "__main__":
    main()
